"""export_comparison_workbook 每物理量分 sheet 工作簿导出测试。

覆盖：每物理量独立 sheet（页首评判标准块 + 活公式明细表）、距离段阈值公式、
仅有效帧导出、不合格红字、汇总页准确率算式/通过条件列与不通过行标红。
"""
import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from radar_wave_analyzer.comparison.exporter import (
    derive_quantity_pairs,
    export_comparison_workbook,
    export_comparison_workbook_bytes,
)

RULES = {
    'accuracy_requirement': 0.9545,
    'three_frame_requirement': 1.0,
    'min_samples_for_conclusion': 3,
    'valid_distance_range': [0, 150],
    'metrics': {
        'Dx': {
            'label': 'Dx', 'unit': 'm', 'mode': 'binned',
            'percent_basis': 'distance',
            'bins': [
                {'range': [0, 10],
                 'normal_limit': {'mode': 'percent', 'percent': 0.02},
                 'three_frame_limit': {'mode': 'absolute', 'absolute': 0.5}},
                {'range': [10, 20],
                 'normal_limit': {'mode': 'absolute', 'absolute': 0.3},
                 'three_frame_limit': {'mode': 'absolute', 'absolute': 0.6}},
            ],
        },
    },
}


def _aligned_df(n: int = 6) -> pd.DataFrame:
    return pd.DataFrame({
        'timestamp': [f't{i}' for i in range(n)],
        'radar_frame': np.arange(1, n + 1),
        'rtk_nearest_ts_parsed': np.arange(1.0, n + 1.0),
        'time_diff_ms': np.full(n, 5.0),
        'radar[Dx]': 5.0 + np.arange(n) * 0.1,
        'rtk[Dx]': np.full(n, 5.0),
    })


def _frames(n: int = 6, abs_errors=None, bin_idx=None, valid_mask=None) -> pd.DataFrame:
    errors = np.full(n, 0.05) if abs_errors is None else np.asarray(abs_errors, float)
    bins = np.zeros(n, dtype=int) if bin_idx is None else np.asarray(bin_idx, int)
    limits = np.where(bins == 0, 0.1, 0.3)
    return pd.DataFrame({
        'radar_frame': np.arange(1, n + 1),
        'timestamp': [f't{i}' for i in range(n)],
        'distance_bin': bins,
        'truth_distance': np.where(bins == 0, 5.0, 15.0),
        'radar_value': 5.0 + errors,
        'truth_value': np.full(n, 5.0),
        'abs_error': errors,
        'normal_limit': limits,
        'normal_pass': errors < limits,
        'three_frame_limit': np.where(bins == 0, 0.5, 0.6),
        'three_frame_violation': np.zeros(n, dtype=bool),
        'valid': np.ones(n, dtype=bool) if valid_mask is None else valid_mask,
    })


def _perf_result(mode: str, frames) -> dict:
    return {
        'available': True, 'mode': mode, 'frames': frames,
        'bins': [
            {'distance_bin': '[0, 10)', 'sample_count': 3, 'rmse': 0.11,
             'normal_pass_count': 2, 'accuracy': 2 / 3, 'status': '不通过'},
            {'distance_bin': '[10, 20)', 'sample_count': 3, 'rmse': 0.42,
             'normal_pass_count': 3, 'accuracy': 1.0, 'status': '通过'},
        ] if mode == 'binned' else [
            {'distance_bin': '完整曲线', 'sample_count': 6, 'rmse': 0.3,
             'status': '通过'},
        ],
        'overall': {}, 'coverage': {},
    }


QUANTITIES = {'Dx': {'label': 'Dx', 'unit': 'm',
                     'radar_col': 'radar[Dx]', 'rtk_col': 'rtk[Dx]'}}


def _wb_bytes(aligned, quantities, perf, rules=RULES) -> bytes:
    return export_comparison_workbook_bytes(aligned, quantities, perf, rules=rules)


def test_per_quantity_sheets_with_formula_columns(tmp_path):
    """每个物理量独立 sheet：页首评判块 + 活公式列（阈值/判定/RMSE/准确率）。"""
    errors = [0.05, 0.20, 0.05, 0.25, 0.05, 0.05]
    perf = {'Dx': _perf_result(
        'binned', _frames(6, abs_errors=errors, bin_idx=[0, 0, 0, 1, 1, 1]))}
    content = _wb_bytes(_aligned_df(6), QUANTITIES, perf)

    path = tmp_path / 'wb.xlsx'
    path.write_bytes(content)
    wb = load_workbook(str(path))

    assert wb.sheetnames == ['验收汇总', 'Dx']
    ws = wb['Dx']
    texts = [str(ws.cell(row=r, column=1).value) for r in range(1, 9)]
    assert any('Dx' in t and '逐帧验收明细' in t for t in texts)
    assert any('准确率' in t and '95.45%' in t for t in texts)

    header_row = next(r for r in range(1, 30)
                      if ws.cell(row=r, column=1).value == '雷达时间戳')
    headers = [ws.cell(row=header_row, column=c).value for c in range(1, 16)]
    assert headers == ['雷达时间戳', '雷达数据', '真值时间戳', '真值数据',
                       '时间差(ms)', '真值距离', '绝对误差', '单帧阈值',
                       '单帧判定', '三帧阈值', '三帧归一化', '三帧违规',
                       '距离段', '段RMSE', '段准确率']
    first = header_row + 1
    # 活公式：绝对误差/阈值（distance 基准引用 F 列）/判定/三帧归一化
    assert ws.cell(row=first, column=7).value == f'=ABS(B{first}-D{first})'
    assert ws.cell(row=first, column=8).value == f'=0.02*ABS(F{first})'
    assert ws.cell(row=first, column=9).value == f'=IF(G{first}<H{first},"✓","✗")'
    assert ws.cell(row=first, column=11).value == f'=IF(J{first}=0,"",G{first}/J{first})'
    # 三帧违规以中间帧归属输出一次：数据首帧无前帧 → 空；帧 2 → 窗口(1,2,3)
    assert ws.cell(row=first, column=12).value is None
    assert (ws.cell(row=first + 1, column=12).value
            == f'=IF(MIN(K{first},K{first + 1},K{first + 2})>=1,"违规","")')
    # 段 RMSE 仅每段首行输出公式；段准确率公式 + 百分比格式
    rmse_formula = ws.cell(row=first, column=14).value
    assert 'SUMPRODUCT' in rmse_formula and 'COUNTIF' in rmse_formula
    assert ws.cell(row=first + 1, column=14).value is None   # 同段第二行留空
    assert ws.cell(row=first + 3, column=14).value is not None  # 新段首行有公式
    acc_formula = ws.cell(row=first, column=15).value
    assert acc_formula.startswith('=COUNTIFS(')
    assert ws.cell(row=first, column=15).number_format == '0.00%'
    # 距离段 1 的阈值为 absolute 常量 0.3
    assert ws.cell(row=first + 3, column=8).value == 0.3
    # 时间戳：雷达原始时间 + 真值 epoch 格式化
    assert ws.cell(row=first, column=1).value == 't0'
    assert ws.cell(row=first, column=3).value.startswith('1970-01-01 00:00:01')


def test_fail_rows_marked_red_in_detail_sheet(tmp_path):
    """单帧判定 ✗ 的行（按后端 normal_pass）：误差/阈值/判定三格红字加粗。"""
    errors = [0.05, 0.20, 0.05, 0.05, 0.05, 0.05]
    perf = {'Dx': _perf_result(
        'binned', _frames(6, abs_errors=errors, bin_idx=[0, 0, 0, 1, 1, 1]))}
    path = tmp_path / 'wb.xlsx'
    path.write_bytes(_wb_bytes(_aligned_df(6), QUANTITIES, perf))

    ws = load_workbook(str(path))['Dx']
    header_row = next(r for r in range(1, 30)
                      if ws.cell(row=r, column=1).value == '雷达时间戳')
    # 判定列读回为公式串（openpyxl 不计算公式），红字按后端判定写入
    assert str(ws.cell(row=header_row + 1, column=9).value).startswith('=IF(')
    # 帧 2（误差 0.20 ≥ 阈值 0.1，normal_pass=False）→ 三格红字加粗
    for col in (7, 8, 9):
        font = ws.cell(row=header_row + 2, column=col).font
        assert str(font.color.rgb).endswith('DC2626')
        assert font.bold
    # 帧 1（合格）保持默认字体
    for col in (7, 8, 9):
        font = ws.cell(row=header_row + 1, column=col).font
        rgb = str(font.color.rgb) if font.color else ''
        assert not rgb.endswith('DC2626')


def test_only_valid_frames_exported(tmp_path):
    """未匹配/段外帧（abs_error 或 normal_limit 非有限值）不进入明细表。"""
    frames = _frames(4, bin_idx=[0, 0, 1, 1])
    frames.loc[2, 'abs_error'] = np.nan      # 未匹配帧
    frames.loc[3, 'normal_limit'] = np.nan   # 距离段外
    perf = {'Dx': _perf_result('binned', frames)}
    path = tmp_path / 'wb.xlsx'
    path.write_bytes(_wb_bytes(_aligned_df(4), QUANTITIES, perf))

    ws = load_workbook(str(path))['Dx']
    header_row = next(r for r in range(1, 30)
                      if ws.cell(row=r, column=1).value == '雷达时间戳')
    data_rows = 0
    r = header_row + 1
    while ws.cell(row=r, column=1).value is not None:
        data_rows += 1
        r += 1
    assert data_rows == 2
    assert ws.cell(row=header_row + 1, column=1).value == 't0'
    assert ws.cell(row=header_row + 2, column=1).value == 't1'


def test_summary_sheet_has_accuracy_formula_columns_and_fail_red(tmp_path):
    """汇总页新增“准确率计算”算式列与“通过条件”列；不通过行红字。"""
    perf = {'Dx': _perf_result(
        'binned', _frames(6, bin_idx=[0, 0, 0, 1, 1, 1]))}
    path = tmp_path / 'wb.xlsx'
    path.write_bytes(_wb_bytes(_aligned_df(6), QUANTITIES, perf))

    ws = load_workbook(str(path))['验收汇总']
    headers = [c.value for c in ws[1]]
    assert '准确率计算' in headers and '通过条件' in headers
    acc_col = headers.index('准确率计算') + 1
    cond_col = headers.index('通过条件') + 1
    status_col = headers.index('结论') + 1
    fail_rows = [row for row in ws.iter_rows(min_row=2)
                 if row[status_col - 1].value == '不通过']
    assert len(fail_rows) == 1
    row = fail_rows[0]
    assert row[acc_col - 1].value == '2/3×100%=66.67%'
    assert row[cond_col - 1].value == '准确率>95.45% 且 连续三帧指标<1'
    for cell in row:
        assert str(cell.font.color.rgb).endswith('DC2626')


def test_without_perf_results(tmp_path):
    """无验收结果：仅汇总提示行，物理量 sheet 写“无逐帧评估数据”。"""
    path = tmp_path / 'wb.xlsx'
    path.write_bytes(_wb_bytes(_aligned_df(3), QUANTITIES, {}))

    wb = load_workbook(str(path))
    assert '无性能评估结果' in wb['验收汇总'].cell(row=1, column=1).value
    texts = [wb['Dx'].cell(row=r, column=1).value for r in range(1, 20)]
    assert any('无逐帧评估数据' in str(t) for t in texts)


def test_empty_aligned_df_returns_hint(tmp_path):
    msg = export_comparison_workbook(
        pd.DataFrame(), QUANTITIES, {}, str(tmp_path / 'wb.xlsx'))
    assert msg == '无对齐数据可导出'


def test_derive_quantity_pairs_prefers_error_column_pairing():
    """error[雷达-真值] 编码配对：真值字段名与雷达字段名不同时也能配对。"""
    aligned = pd.DataFrame({
        'timestamp': ['t1'],
        'radar[Dx]': [1.0],
        'rtk[TruthDx]': [1.0],
        'error[Dx-TruthDx]': [0.0],
        'radar[Ax]': [0.1],
        'rtk[Ax]': [0.1],
        'error[Ax-Ax]': [0.0],
    })
    pairs = derive_quantity_pairs(aligned)
    assert pairs['Dx'] == ('radar[Dx]', 'rtk[TruthDx]')
    assert pairs['Ax'] == ('radar[Ax]', 'rtk[Ax]')
