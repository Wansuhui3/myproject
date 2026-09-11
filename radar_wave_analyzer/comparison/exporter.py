"""
对比结果导出模块（工作簿组装层）。

导出真值对齐与性能验收 xlsx 工作簿（文件路径 / 浏览器下载二进制）。

模块职责拆分（阶段 6）：
  - export_format  格式化辅助：值类型转换、时间戳格式化、限值 Excel 公式、
                   评判标准描述与 sheet 名安全化（零内部依赖）；
  - exporter（本模块）数据对提取、验收汇总构造与工作簿组装写入。
"""
import logging
import os

import numpy as np
import pandas as pd

from ..config import get
from .export_format import (
    _criteria_lines,
    _fmt_epoch_ts,
    _limit_excel,
    _py,
    _sheet_safe_name,
)
from .performance import STATUS_FAIL

logger = logging.getLogger(__name__)


_PERF_SUMMARY_COLUMNS = [
    'metric', 'distance_bin', 'sample_count', 'rmse',
    'normal_pass_count', 'accuracy', 'accuracy_requirement',
    'accuracy_pass', 'three_frame_index', 'three_frame_pass',
    'violation_window_count', 'longest_violation_run',
    'worst_window_start', 'worst_window_end',
    'max_error_in_worst_window', 'status', 'failure_reasons',
]


def _build_performance_summary_df(results: dict) -> pd.DataFrame:
    """构造验收汇总表：每物理量、每距离段一行，供 xlsx「验收汇总」sheet 使用。"""
    rows = []
    for metric, result in (results or {}).items():
        if not isinstance(result, dict):
            continue
        for b in result.get('bins') or []:
            row = {'metric': metric}
            for col in _PERF_SUMMARY_COLUMNS[1:]:
                value = b.get(col)
                if col == 'accuracy' and isinstance(value, (int, float)):
                    value = round(float(value) * 100, 2)
                if col == 'failure_reasons':
                    value = '; '.join(b.get('failure_reasons') or [])
                row[col] = value
            rows.append(row)
    return pd.DataFrame(rows, columns=_PERF_SUMMARY_COLUMNS)


_STANDARD_PAIRS = {
    'radar_Dx': ('Dx', 'rtk_center_x'),
    'radar_Dy': ('Dy', 'rtk_center_y'),
    'radar_Vx': ('Vx', 'rtk_Vx'),
    'radar_Vy': ('Vy', 'rtk_Vy'),
}


def derive_quantity_pairs(aligned_df: pd.DataFrame) -> dict:
    """从对齐结果推导 物理量标签 → (雷达列, 真值列) 配对。

    优先级：
    1. ``error[雷达字段-真值字段]`` 列直接编码了映射配对（真值字段名
       可以与雷达字段名不同）；
    2. 同名直配 ``radar[X]`` ↔ ``rtk[X]``；
    3. 标准通道 ``radar_Dx`` ↔ ``rtk_center_x`` 等（无自定义映射时）。
    """
    columns = [str(c) for c in aligned_df.columns]
    pairs: dict[str, tuple[str, str]] = {}

    def _resolve(label: str, radar_col: str, rtk_col: str) -> None:
        rtk = rtk_col if rtk_col in columns else ''
        if radar_col in columns and label not in pairs:
            pairs[label] = (radar_col, rtk)

    # 1) error[radar-rtk] 编码配对；rtk 字段名可能含 '-'，逐切分点验证
    for col in columns:
        if not col.startswith('error[') or not col.endswith(']'):
            continue
        body = col[len('error['):-1]
        for idx, ch in enumerate(body):
            if ch != '-':
                continue
            radar_label, rtk_label = body[:idx], body[idx + 1:]
            radar_col, rtk_col = f'radar[{radar_label}]', f'rtk[{rtk_label}]'
            if radar_col in columns and rtk_col in columns:
                _resolve(radar_label, radar_col, rtk_col)
                break

    # 2) 同名直配
    for col in columns:
        if col.startswith('radar[') and col.endswith(']'):
            label = col[len('radar['):-1]
            _resolve(label, col, f'rtk[{label}]')

    # 3) 标准通道（无任何 radar[...] 自定义列时兜底）
    if not pairs:
        for radar_col, (label, rtk_col) in _STANDARD_PAIRS.items():
            _resolve(label, radar_col, rtk_col)
    return pairs


def _perf_bin_lookup_maps(perf_result: dict) -> tuple[dict, dict]:
    """根据验收结果构造段索引 → (距离段标签, RMSE) 映射，键为 frames.distance_bin。

    binned 模式 bins[i] 与帧内 bin_index 一一对应；single_limit 模式帧内
    distance_bin 恒为 -1，唯一 bins 行即“完整曲线”。
    """
    bins = perf_result.get('bins') or []
    if perf_result.get('mode') == 'single_limit':
        first = bins[0] if bins else {}
        return ({-1: first.get('distance_bin') or '完整曲线'},
                {-1: first.get('rmse')})
    label_map: dict = {}
    rmse_map: dict = {}
    for index, b in enumerate(bins):
        label_map[index] = b.get('distance_bin')
        rmse_map[index] = b.get('rmse')
    return label_map, rmse_map


def _perf_frame_columns(
        frames: pd.DataFrame, label: str,
        label_map: dict, rmse_map: dict) -> dict:
    """从逐帧结果构造单物理量性能列（绝对误差/阈值/通过/三帧违规/距离段/RMSE）。"""
    index = pd.to_numeric(frames.get('distance_bin'), errors='coerce')
    distance_labels = index.map(
        lambda v: label_map.get(int(v)) if pd.notna(v) and int(v) in label_map else None)
    distance_rmse = index.map(
        lambda v: rmse_map.get(int(v)) if pd.notna(v) and int(v) in rmse_map else None)
    return {
        f'{label}|绝对误差': frames.get('abs_error'),
        f'{label}|普通阈值': frames.get('normal_limit'),
        f'{label}|通过': pd.Series(frames['normal_pass']).map({True: '✓', False: '✗'}),
        f'{label}|三帧违规': pd.Series(frames['three_frame_violation']).map({True: '违规'}),
        f'{label}|距离段': distance_labels,
        f'{label}|距离段RMSE': distance_rmse,
    }


# Sheet2「验收汇总」中文表头（RMSE 为通用缩写保留英文）
_SUMMARY_HEADER_ZH = {
    'metric': '物理量',
    'distance_bin': '距离段',
    'sample_count': '样本数',
    'rmse': 'RMSE',
    'normal_pass_count': '合格帧数',
    'accuracy': '准确率(%)',
    'accuracy_requirement': '准确率要求',
    'accuracy_pass': '准确率通过',
    'three_frame_index': '连续三帧指标',
    'three_frame_pass': '三帧通过',
    'violation_window_count': '违规窗口数',
    'longest_violation_run': '最长连续违规帧',
    'worst_window_start': '最严重窗口起点',
    'worst_window_end': '最严重窗口终点',
    'max_error_in_worst_window': '窗口最大误差',
    'status': '结论',
    'failure_reasons': '不通过原因',
}


def export_comparison_workbook(
        aligned_df: pd.DataFrame,
        quantities: dict,
        perf_results: dict,
        filepath: str,
        rules: dict | None = None,
) -> str:
    """导出真值对齐与性能验收工作簿（xlsx）到文件路径。

    每个物理量一个 sheet（页首评判标准块 + 活公式明细表，不合格红字），
    另有全局「验收汇总」sheet（含准确率算式与通过条件列，不通过行红字）。
    """
    try:
        if aligned_df is None or len(aligned_df) == 0:
            return '无对齐数据可导出'
        parent_dir = os.path.dirname(os.path.abspath(filepath))
        os.makedirs(parent_dir, exist_ok=True)
        # 文件被 Excel 打开时会锁定（Errno 13），自动加时间戳后缀降级保存
        target = filepath
        try:
            with open(target, 'a'):
                pass
        except PermissionError:
            stem, ext = os.path.splitext(filepath)
            target = f'{stem}_{pd.Timestamp.now().strftime("%H%M%S")}{ext}'

        n, sheet_count = _write_workbook(
            aligned_df, quantities, perf_results, target, rules=rules)
        logger.info('导出真值对齐与性能验收工作簿: %s', target)
        return (f'工作簿已导出: {os.path.basename(target)} '
                f'({n}帧, 验收汇总+{sheet_count}个物理量sheet)')
    except Exception as e:
        logger.exception('导出工作簿失败')
        return f'导出失败: {e}'


_DETAIL_HEADERS = [
    '雷达时间戳', '雷达数据', '真值时间戳', '真值数据', '时间差(ms)',
    '真值距离', '绝对误差', '单帧阈值', '单帧判定', '三帧阈值',
    '三帧归一化', '三帧违规', '距离段', '段RMSE', '段准确率',
    '三帧参考线', '帧序号', '不合格误差', '三帧违规点']
_BASIS_COL_DISTANCE = 'F'   # 真值距离列（percent_basis=distance 时阈值公式引用）
_BASIS_COL_TRUTH = 'D'      # 真值数据列（percent_basis=truth 时阈值公式引用）


def _write_workbook(
    aligned_df: pd.DataFrame,
    quantities: dict,
    perf_results: dict,
    target,
    rules: dict | None = None,
) -> tuple[int, int]:
    """写入工作簿到 target（文件路径或二进制缓冲区）。

    每个物理量一个 sheet：页首为评判标准块（判定公式语义与逐段阈值配置），
    明细表含雷达/真值时间戳与数据、活公式（绝对误差、单帧阈值、单帧判定、
    三帧归一化、段RMSE、段准确率），不合格处红色字体；另有全局「验收汇总」
    sheet（含准确率算式与通过条件列，不通过行红字）。
    返回 (对齐帧数, 物理量 sheet 数)。
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    perf_cfg = rules if rules is not None else (get('performance_metrics', {}) or {})
    all_rules = perf_cfg.get('metrics') or {}
    accuracy_req = float(perf_cfg.get('accuracy_requirement', 0.9545))
    three_req = float(perf_cfg.get('three_frame_requirement', 1.0))
    condition_text = (f'准确率>{accuracy_req * 100:.2f}% 且 '
                      f'连续三帧指标<{three_req:g}')
    red_font = Font(color='FFDC2626', bold=True)

    n = len(aligned_df)
    base = aligned_df.reset_index(drop=True)
    rtk_ts_all = base.get('rtk_nearest_ts_parsed')
    time_diff_all = base.get('time_diff_ms')

    wb = Workbook()
    used_names = {'验收汇总'}

    # ── 验收汇总（工作簿首位）──
    ws_sum = wb.active
    ws_sum.title = '验收汇总'
    summary = _build_performance_summary_df(perf_results)
    if summary.empty:
        ws_sum.append(['无性能评估结果（执行对齐后可导出）'])
    else:
        acc_calc, cond = [], []
        for _, row in summary.iterrows():
            sc = int(row.get('sample_count') or 0)
            pc = row.get('normal_pass_count')
            acc = row.get('accuracy')
            if sc > 0 and acc is not None:
                # accuracy 已为百分数（如 66.67）
                acc_calc.append(
                    f'{int(pc) if pc is not None else 0}/{sc}×100%={acc:.2f}%')
                cond.append(condition_text)
            else:
                acc_calc.append('—')
                cond.append('—')
        summary = summary.rename(columns=_SUMMARY_HEADER_ZH)
        summary['准确率计算'] = acc_calc
        summary['通过条件'] = cond
        ws_sum.append(list(summary.columns))
        for _, row in summary.iterrows():
            ws_sum.append([_py(v) for v in row.tolist()])
        ws_sum.freeze_panes = 'A2'
        ws_sum.auto_filter.ref = ws_sum.dimensions
        status_col = summary.columns.get_loc('结论')
        for row in ws_sum.iter_rows(min_row=2):
            if row[status_col].value == STATUS_FAIL:
                for cell in row:
                    cell.font = red_font
        for col_idx, col_name in enumerate(summary.columns, start=1):
            ws_sum.column_dimensions[get_column_letter(col_idx)].width = (
                24 if col_name in ('准确率计算', '通过条件', '不通过原因')
                else 12 if col_name in ('物理量', '结论', 'RMSE') else 13)


    # ── 每物理量一个 sheet ──
    for label, qty in (quantities or {}).items():
        ws = wb.create_sheet(_sheet_safe_name(label, used_names))
        rule = all_rules.get(label) or {}
        unit = str(qty.get('unit') or rule.get('unit') or '')
        perf_result = (perf_results or {}).get(label) or {}
        frames = perf_result.get('frames')
        basis_col = ('_BASIS_COL_DISTANCE'
                     if str(rule.get('percent_basis', 'truth')) == 'distance'
                     else '_BASIS_COL_TRUTH')
        basis_col = _BASIS_COL_DISTANCE if (
            str(rule.get('percent_basis', 'truth')) == 'distance'
        ) else _BASIS_COL_TRUTH

        for line in _criteria_lines(label, unit, rule, perf_cfg):
            ws.append([line])
        ws.append([None])  # 空行（append([]) 不占行，会导致行号错位）

        if not isinstance(frames, pd.DataFrame) or len(frames) == 0:
            ws.append(['无逐帧评估数据（执行对齐并产生有效样本后可导出）'])
            continue

        fr = frames.reset_index(drop=True)
        bin_idx = pd.to_numeric(fr.get('distance_bin'), errors='coerce')
        abs_err = pd.to_numeric(fr.get('abs_error'), errors='coerce').to_numpy(
            dtype=float)
        limits = pd.to_numeric(fr.get('normal_limit'), errors='coerce').to_numpy(
            dtype=float)
        mask = np.isfinite(abs_err) & np.isfinite(limits)
        is_binned = perf_result.get('mode') != 'single_limit'
        if is_binned:
            mask &= (bin_idx.fillna(-1).to_numpy(dtype=int) >= 0)
        if not mask.any():
            ws.append(['无有效匹配帧'])
            continue

        label_map, _rmse_map = _perf_bin_lookup_maps(perf_result)
        rule_bins = rule.get('bins') or []
        same_len = len(fr) == n
        row_positions = np.flatnonzero(mask)
        header_row = ws.max_row + 1
        ws.append(_DETAIL_HEADERS)
        first = header_row + 1
        last = header_row + len(row_positions)
        breaks_series = (
            pd.to_numeric(fr.get('continuity_break'), errors='coerce').fillna(0)
            .to_numpy(dtype=int) if 'continuity_break' in fr else None)

        # ── 预取列数据到 Python 列表 ──
        # 明细循环逐帧执行，循环内 Series.iloc 标量索引的 pandas 开销会随
        # 行数线性放大（实测为生成耗时主因之一）；预转列表后循环仅做下标访问。
        def _rows(name):
            return fr[name].tolist() if name in fr else None

        ts_list = _rows('timestamp')
        radar_list = _rows('radar_value')
        truth_list = _rows('truth_value')
        distance_list = _rows('truth_distance')
        three_limit_list = _rows('three_frame_limit')
        normal_pass_list = _rows('normal_pass')
        violation_list = _rows('three_frame_violation')
        rtk_ts_list = (rtk_ts_all.tolist()
                       if rtk_ts_all is not None and same_len else None)
        time_diff_list = (time_diff_all.tolist()
                          if time_diff_all is not None and same_len else None)
        bin_idx_list = bin_idx.tolist()

        prev_bi = None
        fail_rows: list[int] = []
        violation_rows: list[int] = []
        for offset, i in enumerate(row_positions):
            r = first + offset
            bi_value = bin_idx_list[i]
            bi = int(bi_value) if pd.notna(bi_value) else -1
            limit_rule = (rule_bins[bi].get('normal_limit')
                          if is_binned and 0 <= bi < len(rule_bins) else None)
            threshold = _limit_excel(limit_rule, basis_col, r)
            if threshold is None:
                threshold = _py(limits[i])
            violation = (bool(violation_list[i])
                         if violation_list is not None else False)
            normal_pass = (bool(normal_pass_list[i])
                           if normal_pass_list is not None else True)
            is_break = bool(breaks_series[i]) if breaks_series is not None else False
            if is_binned:
                dist_label = label_map.get(bi)
            else:
                dist_label = label_map.get(-1, '完整曲线')
            rtk_ts = (_fmt_epoch_ts(rtk_ts_list[i])
                      if rtk_ts_list is not None else '')
            time_diff = (_py(time_diff_list[i])
                         if time_diff_list is not None else None)

            # ── 三帧违规（Excel 活公式，binned 模式）──
            # 违规 = 连续三帧归一化最小值 MIN(K前, K本, K后) ≥ 1，内联公式，
            # 以中间帧归属输出一次：
            #   段末帧（下一行中断或最后一行）：无完整三帧窗口 → 空
            #   连续段首帧（数据首行或中断后首帧）：无前帧 → 空
            #   其余：MIN(K-1, K, K+1)
            if is_binned:
                next_is_break = (
                    breaks_series is not None
                    and offset + 1 < len(row_positions)
                    and bool(breaks_series[row_positions[offset + 1]]))
                is_seg_tail = (r == last) or next_is_break
                is_seg_head = (r == first) or is_break
                if is_seg_tail or is_seg_head:
                    violation_formula = None
                else:
                    violation_formula = (
                        f'=IF(MIN(K{r - 1},K{r},K{r + 1})>=1,"违规","")')
            else:
                violation_formula = None

            # ── 段RMSE：每段只在首行输出一个公式，其余行留空 ──
            is_bin_head = (offset == 0) or is_break or (bi != prev_bi)
            rmse_formula = None
            if is_binned and is_bin_head:
                rmse_formula = (
                    f'=SQRT(SUMPRODUCT(($M${first}:$M${last}=M{r})'
                    f'*($G${first}:$G${last})^2)/COUNTIF($M${first}:$M${last},M{r}))')

            ws.append([
                str(ts_list[i]) if ts_list is not None else '',
                _py(radar_list[i]) if radar_list is not None else None,
                rtk_ts,
                _py(truth_list[i]) if truth_list is not None else None,
                time_diff,
                _py(distance_list[i]) if distance_list is not None else None,
                f'=ABS(B{r}-D{r})',
                threshold,
                f'=IF(G{r}<H{r},"✓","✗")',
                _py(three_limit_list[i]) if three_limit_list is not None else None,
                f'=IF(J{r}=0,"",G{r}/J{r})',
                violation_formula,
                dist_label,
                rmse_formula,
                (f'=COUNTIFS($M${first}:$M${last},M{r},'
                 f'$I${first}:$I${last},"✓")/COUNTIF($M${first}:$M${last},M{r})'),
            ])
            # 不合格帧：误差/阈值/判定三格红字；违规帧：三帧违规格红字
            # （按后端判定结果，公式列保存后读不到计算值，无法事后判断）
            if not normal_pass:
                fail_rows.append(r)
            if violation:
                violation_rows.append(r)
            prev_bi = bi

        ws.freeze_panes = f'A{first}'
        ws.auto_filter.ref = f'A{header_row}:O{last}'
        # 着色与数字格式在逐帧循环结束后统一处理，避免热循环内反复查找单元格
        for r in fail_rows:
            for col in (7, 8, 9):
                ws.cell(row=r, column=col).font = red_font
        for r in violation_rows:
            ws.cell(row=r, column=12).font = red_font
        for (cell,) in ws.iter_rows(min_row=first, max_row=last,
                                    min_col=15, max_col=15):
            cell.number_format = '0.00%'
        widths = [22, 12, 22, 12, 10, 10, 10, 12, 9, 10, 11, 9, 12, 12, 10]
        for col_idx, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = width

    wb.save(target)
    return n, max(len(quantities or {}), 0)


def export_comparison_workbook_bytes(
        aligned_df: pd.DataFrame,
        quantities: dict,
        perf_results: dict,
        rules: dict | None = None,
) -> bytes:
    """在内存中生成工作簿，返回 xlsx 二进制内容（供浏览器下载）。"""
    import io

    buf = io.BytesIO()
    _write_workbook(aligned_df, quantities, perf_results, buf, rules=rules)
    return buf.getvalue()
